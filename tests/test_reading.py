import io
from pathlib import Path

import pytest

from novel_ai.reading import extract_reference_text


def test_txt_is_decoded_directly():
    data = "他推开门。屋里没人。".encode("utf-8")
    assert extract_reference_text("ref.txt", data) == "他推开门。屋里没人。"


def test_docx_paragraphs_are_preserved():
    from docx import Document

    doc = Document()
    doc.add_paragraph("第一段：夜班电台。")
    doc.add_paragraph("")
    doc.add_paragraph("第二段：23:17，电话准时打入。")
    buf = io.BytesIO()
    doc.save(buf)
    text = extract_reference_text("ref.docx", buf.getvalue())
    assert text.startswith("第一段：夜班电台。")
    assert "第二段：23:17，电话准时打入。" in text
    assert "\n\n\n" not in text


def make_pdf(text: str) -> bytes:
    """Hand-build a minimal valid single-page PDF so tests need no PDF writer."""
    stream = f"BT /F1 18 Tf 72 720 Td ({text}) Tj ET".encode("latin-1")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, obj in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + obj + b"\nendobj\n"
    xref_pos = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode()
    out += b"0000000000 65535 f \n"
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_pos}\n%%EOF"
    ).encode()
    return bytes(out)


def test_pdf_text_is_extracted():
    text = extract_reference_text("ref.pdf", make_pdf("Hello PDF extraction"))
    assert "Hello PDF extraction" in text


def test_epub_without_optional_reader_raises_clear_dependency_error():
    with pytest.raises(ValueError, match="Docling|MarkItDown"):
        extract_reference_text("ref.epub", b"whatever")


def test_corrupt_pdf_raises_pdf_error_not_silent_empty():
    from pypdf.errors import PdfReadError

    with pytest.raises((PdfReadError, ValueError)):
        extract_reference_text("bad.pdf", b"not a pdf at all")


def test_reference_text_strips_invisible_prompt_injection_controls():
    data = "正常\u200b文本\u202e倒序".encode("utf-8")
    assert extract_reference_text("ref.txt", data) == "正常文本倒序"


def _trafilatura_only(monkeypatch):
    from novel_ai import reading

    real_find_spec = reading.importlib.util.find_spec

    def find_spec(name):
        if name in {"docling", "markitdown"}:
            return None
        return real_find_spec(name)

    monkeypatch.setattr(reading.importlib.util, "find_spec", find_spec)


def test_trafilatura_extracts_uploaded_html_fixture_and_preserves_cleaning(monkeypatch):
    pytest.importorskip("trafilatura")
    _trafilatura_only(monkeypatch)
    fixture = (Path(__file__).parent / "fixtures" / "reference.html").read_bytes()

    from novel_ai.reading import extract_reference

    result = extract_reference(
        "reference.html",
        fixture,
        normalization_map={"繁": "简"},
    )

    assert result.backend == "trafilatura"
    assert "正文标记：夜班电台的门在23:17准时响起" in result.text
    assert "导航不应出现" not in result.text
    assert "页脚不应出现" not in result.text
    assert "评论不应出现" not in result.text
    assert "表格数据不应出现" not in result.text
    assert "零宽字符" in result.text
    assert "繁" not in result.text


def test_trafilatura_empty_or_bad_html_has_clear_fallback_error(monkeypatch):
    pytest.importorskip("trafilatura")
    _trafilatura_only(monkeypatch)

    from novel_ai.reading import extract_reference

    for payload in (
        b"<html><head><title></title></head><body><div></div></body></html>",
        b"",
    ):
        with pytest.raises(ValueError, match="trafilatura"):
            extract_reference("empty.html", payload)


def test_html_without_optional_backend_keeps_dependency_fallback(monkeypatch):
    from novel_ai import reading

    monkeypatch.setattr(reading.importlib.util, "find_spec", lambda name: None)

    from novel_ai.reading import extract_reference

    with pytest.raises(ValueError, match="Docling|MarkItDown"):
        extract_reference("reference.html", "<html><body>正文</body></html>".encode("utf-8"))


def test_trafilatura_html_path_spies_real_extract_and_does_not_open_network(monkeypatch):
    trafilatura = pytest.importorskip("trafilatura")
    _trafilatura_only(monkeypatch)
    import socket

    fixture = (Path(__file__).parent / "fixtures" / "reference.html").read_bytes()
    original_extract = trafilatura.extract
    observed = {}

    def spy_extract(filecontent, *args, **kwargs):
        observed["filecontent"] = filecontent
        observed["kwargs"] = kwargs
        return original_extract(filecontent, *args, **kwargs)

    def blocked_connect(*args, **kwargs):
        raise AssertionError("local HTML extraction attempted a network connection")

    monkeypatch.setattr(trafilatura, "extract", spy_extract)
    monkeypatch.setattr(socket.socket, "connect", blocked_connect)

    from novel_ai.reading import extract_reference

    result = extract_reference("reference.html", fixture)
    assert result.backend == "trafilatura"
    assert observed["filecontent"] == fixture
    assert observed["kwargs"]["url"] is None
    assert observed["kwargs"]["favor_precision"] is True
    assert observed["kwargs"]["include_comments"] is False
    assert observed["kwargs"]["include_tables"] is False


@pytest.mark.parametrize("failure", ["empty", "error"])
def test_trafilatura_empty_or_error_falls_through_to_markitdown(monkeypatch, failure):
    trafilatura = pytest.importorskip("trafilatura")
    from novel_ai import reading
    import sys
    import types

    fixture = (Path(__file__).parent / "fixtures" / "reference.html").read_bytes()
    real_find_spec = reading.importlib.util.find_spec

    def find_spec(name):
        if name == "docling":
            return None
        if name == "markitdown":
            return object()
        return real_find_spec(name)

    monkeypatch.setattr(reading.importlib.util, "find_spec", find_spec)

    fallback_module = types.ModuleType("markitdown")

    class FallbackResult:
        markdown = "MarkItDown fallback text"

    class FallbackMarkItDown:
        def convert(self, path):
            return FallbackResult()

    fallback_module.MarkItDown = FallbackMarkItDown
    monkeypatch.setitem(sys.modules, "markitdown", fallback_module)

    def failing_extract(filecontent, *args, **kwargs):
        assert filecontent == fixture
        if failure == "error":
            raise RuntimeError("synthetic trafilatura failure")
        return None

    monkeypatch.setattr(trafilatura, "extract", failing_extract)

    from novel_ai.reading import extract_reference

    result = extract_reference("reference.html", fixture)
    assert result.backend == "markitdown"
    assert result.text == "MarkItDown fallback text"


